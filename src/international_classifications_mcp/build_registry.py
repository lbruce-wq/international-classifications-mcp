from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
from datetime import UTC, datetime
from html import unescape
from pathlib import Path

from openpyxl import load_workbook

from .catalog import CLASSIFICATIONS, SEED_CODES

PACKAGE = Path(__file__).resolve().parent
PROJECT = PACKAGE.parents[1]
RAW = PROJECT / "data" / "raw"
OUTPUT = PACKAGE / "data" / "registry.sqlite"


SCHEMA = """
PRAGMA journal_mode=DELETE;
CREATE TABLE classifications (
  id TEXT PRIMARY KEY, acronym TEXT NOT NULL, name TEXT NOT NULL, version TEXT NOT NULL,
  custodian TEXT NOT NULL, domain TEXT NOT NULL, status TEXT NOT NULL, coverage TEXT NOT NULL,
  description TEXT NOT NULL, source_url TEXT NOT NULL, licence_note TEXT NOT NULL,
  retrieved_at TEXT, source_sha256 TEXT
);
CREATE TABLE codes (
  classification_id TEXT NOT NULL REFERENCES classifications(id), code TEXT NOT NULL,
  label TEXT NOT NULL, level INTEGER, parent_code TEXT, definition TEXT, includes TEXT,
  excludes TEXT, language TEXT NOT NULL DEFAULT 'en', codelist_id TEXT,
  PRIMARY KEY(classification_id, code, language)
);
CREATE INDEX codes_classification_idx ON codes(classification_id, code);
CREATE VIRTUAL TABLE codes_fts USING fts5(
  classification_id UNINDEXED, code, label, definition, includes, excludes,
  content='codes', content_rowid='rowid', tokenize='unicode61 remove_diacritics 2'
);
CREATE TRIGGER codes_ai AFTER INSERT ON codes BEGIN
  INSERT INTO codes_fts(rowid, classification_id, code, label, definition, includes, excludes)
  VALUES (new.rowid, new.classification_id, new.code, new.label, new.definition, new.includes, new.excludes);
END;
CREATE TABLE correspondences (
  source_classification_id TEXT NOT NULL, source_code TEXT NOT NULL,
  target_classification_id TEXT NOT NULL, target_code TEXT NOT NULL,
  relationship TEXT NOT NULL, official INTEGER NOT NULL DEFAULT 1, note TEXT,
  source_url TEXT NOT NULL
);
CREATE INDEX correspondence_idx ON correspondences(source_classification_id, target_classification_id, source_code);
CREATE TABLE codelists (
  classification_id TEXT NOT NULL REFERENCES classifications(id), codelist_id TEXT NOT NULL,
  title TEXT NOT NULL, concept TEXT NOT NULL, source_version TEXT NOT NULL,
  warning TEXT NOT NULL, source_url TEXT NOT NULL,
  PRIMARY KEY(classification_id, codelist_id)
);
CREATE TABLE build_metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL);
"""


def _clean(value: object | None) -> str | None:
    if value is None:
        return None
    text = str(value).replace("_x000D_", "").replace("\xa0", " ").strip()
    return text or None


def _parent_from_code(code: str, dotted: bool = False) -> str | None:
    if dotted:
        bits = code.split(".")
        return ".".join(bits[:-1]) or None
    if len(code) <= 1:
        return None
    return code[:-1]


def _insert_code(
    conn: sqlite3.Connection, cid: str, code: object, label: object, **kwargs: object
) -> None:
    c, lab = _clean(code), _clean(label)
    if not c or not lab:
        return
    conn.execute(
        "INSERT OR REPLACE INTO codes(classification_id,code,label,level,parent_code,definition,includes,excludes,language,codelist_id) VALUES(?,?,?,?,?,?,?,?,?,?)",
        (
            cid,
            c,
            lab,
            kwargs.get("level"),
            kwargs.get("parent_code"),
            _clean(kwargs.get("definition")),
            _clean(kwargs.get("includes")),
            _clean(kwargs.get("excludes")),
            kwargs.get("language", "en"),
            kwargs.get("codelist_id"),
        ),
    )


def _ingest_simple_csv(conn: sqlite3.Connection, cid: str, filename: str) -> None:
    with (RAW / filename).open(encoding="cp1252", newline="") as handle:
        rows = csv.reader(handle)
        next(rows)
        for code, label, *_ in rows:
            code = code.strip()
            _insert_code(
                conn, cid, code, label, level=len(code), parent_code=_parent_from_code(code)
            )


def _ingest_coicop(conn: sqlite3.Connection) -> None:
    ws = load_workbook(RAW / "coicop2018.xlsx", read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    next(rows)
    for code, title, intro, includes, also_includes, excludes, *_ in rows:
        if not code:
            continue
        code = str(code)
        combined = "\n".join(filter(None, [_clean(includes), _clean(also_includes)])) or None
        _insert_code(
            conn,
            "coicop2018",
            code,
            title,
            level=code.count(".") + 1,
            parent_code=_parent_from_code(code, True),
            definition=intro,
            includes=combined,
            excludes=excludes,
        )


def _ingest_unsd_text(conn: sqlite3.Connection, cid: str, filename: str) -> None:
    for line in (RAW / filename).read_text(encoding="utf-8-sig", errors="replace").splitlines()[1:]:
        match = re.match(r"^\s*([0-9]+(?:\.[0-9]+)*)\s+(.+?)\s*$", line)
        if not match:
            continue
        code, label = match.groups()
        parent = code.rsplit(".", 1)[0] if "." in code else None
        _insert_code(conn, cid, code, label, level=code.count(".") + 1, parent_code=parent)


def _ingest_icatus(conn: sqlite3.Connection) -> None:
    with (RAW / "icatus2016.txt").open(encoding="cp1252", newline="") as handle:
        for row in csv.DictReader(handle):
            code, label = row["Code"].strip(), row["Description"].strip()
            parent = code[:-1] if len(code) > 2 else None
            _insert_code(conn, "icatus2016", code, label, level=len(code), parent_code=parent)


def _ingest_unece(conn: sqlite3.Connection, cid: str, filename: str, code_field: str) -> None:
    with (RAW / filename).open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            code = (row.get(code_field) or "").strip()
            label = (row.get("Name") or "").strip()
            if code and label:
                _insert_code(conn, cid, code, label, level=1, definition=row.get("Description"))


def _ingest_iscedf(conn: sqlite3.Connection) -> None:
    text = (RAW / "iscedf2013.ttl").read_text(encoding="utf-8")
    for block in re.split(r"\n(?=:[A-Za-z0-9_]+ a skos:Concept)", text):
        labels = re.findall(r'skos:prefLabel "([^"]+)"@en', block)
        codes = re.findall(r'skos:notation "(\d{2,4})"', block)
        if not labels or not codes:
            continue
        code = max(codes, key=len)
        parent = code[:-1] if len(code) in {3, 4} else None
        notes = re.findall(r'skos:scopeNote "([^"]+)"@en', block)
        _insert_code(conn, "iscedf2013", code, labels[0], level=len(code), parent_code=parent, definition=notes[0] if notes else None)


def _ingest_iccs(conn: sqlite3.Connection) -> None:
    text = (RAW / "iccs1.pdf.txt").read_text(encoding="utf-8")[81053:260000]
    for code, label in re.findall(r"(?m)^\s*(\d{4,6})\s+(.\S.*)$", text):
        parent = code[:-1] if len(code) in {5, 6} else code[:2]
        _insert_code(conn, "iccs1", code, label.strip(), level={4: 2, 5: 3, 6: 4}[len(code)], parent_code=parent)


def _ingest_icc(conn: sqlite3.Connection) -> None:
    lines = [line.strip() for line in (RAW / "icc11.txt").read_text(encoding="utf-8-sig", errors="replace").splitlines()]
    start = lines.index("Indicative Crop Classification Version 1.1 (ICC)")
    ignored = {"T", "P", "T/P", "Group", "Class", "Sub-", "class", "Order", "Title", "Crop type*"}
    for index, code in enumerate(lines[start + 1 :], start + 1):
        if not re.fullmatch(r"\d+(?:\.\d+){0,3}", code):
            continue
        label = next((candidate for candidate in lines[index + 1 : index + 10] if candidate and candidate not in ignored and not re.fullmatch(r"\d+(?:\.\d+){0,3}", candidate)), None)
        if not label:
            continue
        parent = code.rsplit(".", 1)[0] if "." in code else None
        _insert_code(conn, "icc11", code, label, level=code.count(".") + 1, parent_code=parent)


def _ingest_curated_codelists(conn: sqlite3.Connection) -> None:
    payload = json.loads((RAW / "survey_codelists_curated.json").read_text(encoding="utf-8"))
    warning = "Use only within the named source/version and verify exact wording, skips and country adaptation against the official instrument."
    for cid, family in payload.items():
        for codelist_id, (title, options) in family["lists"].items():
            conn.execute("INSERT INTO codelists VALUES(?,?,?,?,?,?,?)", (cid, codelist_id, title, title, family["version"], warning, family["source_url"]))
            for code, label in options:
                full_code = f"{codelist_id}.{code}"
                _insert_code(conn, cid, full_code, label, level=2, parent_code=codelist_id, codelist_id=codelist_id, includes=f"Codelist: {title}")


def _ingest_comtrade(conn: sqlite3.Connection, cid: str, filename: str) -> None:
    payload = json.loads((RAW / filename).read_text(encoding="utf-8"))
    for item in payload["results"]:
        code = str(item["id"])
        if code == "TOTAL":
            continue
        label = re.sub(rf"^{re.escape(code)}\s*-\s*", "", str(item["text"])).strip()
        parent = item.get("parent")
        _insert_code(
            conn,
            cid,
            code,
            label,
            level=int(item.get("aggrlevel") or len(code)),
            parent_code=None if parent in (None, "#", "TOTAL") else str(parent),
        )


def _ingest_isic_mapping(conn: sqlite3.Connection) -> None:
    path = RAW / "isic4to5.xlsx"
    ws = load_workbook(path, read_only=True, data_only=True)["ISIC4-5"]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    source_url = next(x["source_url"] for x in CLASSIFICATIONS if x["id"] == "isic4")
    for row in rows:
        if len(row) < 6 or not row[1] or not row[4]:
            continue
        source_code, source_label, target_code = str(row[1]), row[2], str(row[4])
        _insert_code(
            conn,
            "isic4",
            source_code,
            source_label,
            level=len(source_code),
            parent_code=_parent_from_code(source_code),
        )
        change = _clean(row[6]) if len(row) > 6 else None
        note = _clean(row[7]) if len(row) > 7 else None
        relationship = "exact" if not change else "official_correspondence"
        conn.execute(
            "INSERT INTO correspondences VALUES(?,?,?,?,?,?,?,?)",
            (
                "isic4",
                source_code,
                "isic5",
                target_code,
                relationship,
                1,
                note or change,
                source_url,
            ),
        )


def _ingest_isco(conn: sqlite3.Connection) -> None:
    for line in (RAW / "isco08.txt").read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        major = re.fullmatch(r"MAJOR GROUP\s+([0-9]):\s*(.+)", line, flags=re.IGNORECASE)
        if major:
            _insert_code(conn, "isco08", major.group(1), major.group(2), level=1)
            continue
        item = re.fullmatch(r"([0-9]{2,4})\s+(.+)", line)
        if item:
            code, label = item.groups()
            _insert_code(
                conn,
                "isco08",
                code,
                label,
                level=len(code),
                parent_code=_parent_from_code(code),
            )


def _strip_html(value: str) -> str:
    return unescape(re.sub(r"<[^>]+>", "", value)).strip()


def _ingest_m49(conn: sqlite3.Connection) -> None:
    html = (RAW / "m49.csv").read_text(encoding="utf-8", errors="replace")
    table = re.search(r'<table id\s*=\s*"downloadTableEN".*?</table>', html, flags=re.DOTALL)
    if not table:
        raise ValueError("Official M49 English table not found")
    seen: set[str] = set()
    for row in re.findall(r"<tr>(.*?)</tr>", table.group(0), flags=re.DOTALL):
        cells = [
            _strip_html(cell) for cell in re.findall(r"<td[^>]*>(.*?)</td>", row, flags=re.DOTALL)
        ]
        if len(cells) < 12 or not re.fullmatch(r"\d{3}", cells[9] or ""):
            continue
        country, m49, iso2, iso3 = cells[8], cells[9], cells[10], cells[11]
        if m49 in seen:
            continue
        seen.add(m49)
        definition = f"ISO alpha-2: {iso2}; ISO alpha-3: {iso3}; UN region: {cells[3]}; sub-region: {cells[5]}"
        _insert_code(conn, "m49", m49, country, level=3, definition=definition)


def _ingest_mics7(conn: sqlite3.Connection) -> None:
    """Load only round-specific MICS questionnaire response codelists."""
    payload = json.loads((RAW / "mics7_curated.json").read_text(encoding="utf-8"))
    list_titles = {
        "SEX": "Sex", "YN": "Yes, no and permitted non-response", "WS_SOURCE": "Main drinking-water source",
        "WS_SAN": "Sanitation facility type", "BR.STATUS": "Birth-registration status",
        "IM.SOURCE": "Vaccination evidence/source", "CF.DIFFICULTY": "Child-functioning difficulty scale",
        "CD.METHOD_GROUP": "Child-discipline method group", "ED.ATTEND": "School-attendance status",
    }
    list_aliases = {
        "Sex": "SEX", "Yes/no": "YN", "Main drinking-water source": "WS_SOURCE",
        "Sanitation facility": "WS_SAN", "Birth registration": "BR.STATUS",
        "Vaccination evidence": "IM.SOURCE", "Child functioning response scale": "CF.DIFFICULTY",
        "Child discipline method group": "CD.METHOD_GROUP", "School attendance status": "ED.ATTEND",
    }
    provenance_url = "https://classifications.impactengines.ai/provenance/mics7"
    warning = "Verify wording, codes, skips and country customisation against the cited MICS7 questionnaire before deployment."
    for codelist_id, title in list_titles.items():
        conn.execute("INSERT INTO codelists VALUES(?,?,?,?,?,?,?)", ("mics7_responses", codelist_id, title, title, payload["versions"]["responses"], warning, provenance_url))
    for code, label, codelist, definition in payload["responses"]:
        codelist_id = list_aliases[codelist]
        _insert_code(
            conn,
            "mics7_responses",
            code,
            label,
            level=2,
            parent_code=codelist_id,
            codelist_id=codelist_id,
            definition=definition,
            includes=f"Codelist: {codelist}",
        )


def _source_hash(cid: str) -> str | None:
    files = {
        "isic5": "isic5.csv",
        "cpc3": "cpc3.csv",
        "coicop2018": "coicop2018.xlsx",
        "hs2022": "hs2022.json",
        "sitc4": "sitc4.json",
        "bec5": "bec5.json",
        "isic4": "isic4to5.xlsx",
        "isco08": "isco08.txt",
        "m49": "m49.csv",
        "mics7_responses": "mics7_curated.json",
        "icatus2016": "icatus2016.txt", "cofog1999": "cofog1999.txt", "copni1999": "copni1999.txt", "copp1999": "copp1999.txt",
        "unece_rec20": "unece_rec20.csv", "unece_rec21": "unece_rec21.csv", "iscedf2013": "iscedf2013.ttl", "iccs1": "iccs1.pdf",
        "dhs8_responses": "survey_codelists_curated.json", "wg_responses": "survey_codelists_curated.json", "jmp2018_responses": "survey_codelists_curated.json", "who_vax_responses": "survey_codelists_curated.json", "fao_wca2020_responses": "survey_codelists_curated.json", "icc11": "icc11.doc",
    }
    filename = files.get(cid)
    if not filename or not (RAW / filename).exists():
        return None
    return hashlib.sha256((RAW / filename).read_bytes()).hexdigest()


def build(output: Path = OUTPUT) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()
    conn = sqlite3.connect(output)
    conn.executescript(SCHEMA)
    now = datetime.now(UTC).isoformat()
    for item in CLASSIFICATIONS:
        conn.execute(
            "INSERT INTO classifications VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (*item.values(), now, _source_hash(item["id"])),
        )
    for cid, rows in SEED_CODES.items():
        for code, label in rows:
            _insert_code(
                conn, cid, code, label, level=len(code), parent_code=_parent_from_code(code)
            )
    _ingest_simple_csv(conn, "isic5", "isic5.csv")
    _ingest_simple_csv(conn, "cpc3", "cpc3.csv")
    _ingest_coicop(conn)
    _ingest_comtrade(conn, "hs2022", "hs2022.json")
    _ingest_comtrade(conn, "sitc4", "sitc4.json")
    _ingest_comtrade(conn, "bec5", "bec5.json")
    _ingest_isic_mapping(conn)
    _ingest_isco(conn)
    _ingest_m49(conn)
    _ingest_mics7(conn)
    _ingest_unsd_text(conn, "cofog1999", "cofog1999.txt")
    _ingest_unsd_text(conn, "copni1999", "copni1999.txt")
    _ingest_unsd_text(conn, "copp1999", "copp1999.txt")
    _ingest_icatus(conn)
    _ingest_unece(conn, "unece_rec20", "unece_rec20.csv", "CommonCode")
    _ingest_unece(conn, "unece_rec21", "unece_rec21.csv", "Code")
    _ingest_iscedf(conn)
    _ingest_iccs(conn)
    _ingest_icc(conn)
    _ingest_curated_codelists(conn)
    conn.execute("INSERT INTO build_metadata VALUES('registry_version','0.5.1')")
    conn.execute("INSERT INTO build_metadata VALUES('built_at',?)", (now,))
    conn.execute("INSERT INTO build_metadata VALUES('national_census_geography','out_of_scope')")
    conn.commit()
    conn.execute("ANALYZE")
    conn.execute("VACUUM")
    conn.close()
    return output


def main() -> None:
    path = build()
    print(path)


if __name__ == "__main__":
    main()
